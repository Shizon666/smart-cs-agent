# -*- coding: utf-8 -*-
"""多格式文档加载：对齐教程 asset/load 常见类型。"""
from __future__ import annotations

import json
from pathlib import Path

from langchain_core.documents import Document
from langchain_community.document_loaders import (
    TextLoader,
    CSVLoader,
    PyPDFLoader,
)

from rag import config

# 参考 langchain1.2_tutorial/asset/load：txt/csv/json/pdf/docx/md/html
_SUPPORTED = {
    ".txt",
    ".md",
    ".markdown",
    ".html",
    ".htm",
    ".docx",
    ".pdf",
    ".csv",
    ".json",
}


def _meta(path: Path, extra: dict | None = None) -> dict:
    m = {
        "source": str(path),
        "filename": path.name,
        "filetype": path.suffix.lower().lstrip("."),
    }
    if extra:
        m.update(extra)
    return m


def _load_docx(path: Path) -> list[Document]:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return [Document(page_content=text, metadata=_meta(path))]


def _load_json(path: Path) -> list[Document]:
    """支持：list[{text|content|q/a}] 或 {items:[...]} / {messages:[...]}。"""
    data = json.loads(path.read_text(encoding="utf-8"))
    docs: list[Document] = []

    def one(obj: dict | str, idx: int) -> Document | None:
        if isinstance(obj, str):
            content = obj.strip()
        elif isinstance(obj, dict):
            if "text" in obj:
                content = str(obj["text"])
            elif "content" in obj:
                content = str(obj["content"])
            elif "q" in obj and "a" in obj:
                content = f"问：{obj['q']}\n答：{obj['a']}"
            elif "question" in obj and "answer" in obj:
                content = f"问：{obj['question']}\n答：{obj['answer']}"
            else:
                content = json.dumps(obj, ensure_ascii=False)
        else:
            return None
        if not content.strip():
            return None
        return Document(page_content=content, metadata=_meta(path, {"row": idx}))

    if isinstance(data, list):
        for i, item in enumerate(data):
            d = one(item, i)
            if d:
                docs.append(d)
    elif isinstance(data, dict):
        for key in ("items", "faqs", "messages", "data", "documents"):
            if isinstance(data.get(key), list):
                for i, item in enumerate(data[key]):
                    d = one(item, i)
                    if d:
                        docs.append(d)
                break
        else:
            docs.append(
                Document(
                    page_content=json.dumps(data, ensure_ascii=False, indent=2),
                    metadata=_meta(path),
                )
            )
    return docs


def _load_one(path: Path) -> list[Document]:
    suffix = path.suffix.lower()

    if suffix == ".docx":
        return _load_docx(path)
    if suffix == ".pdf":
        docs = PyPDFLoader(str(path)).load()
        for d in docs:
            d.metadata = {**(d.metadata or {}), **_meta(path)}
        return docs
    if suffix == ".csv":
        docs = CSVLoader(str(path), encoding="utf-8").load()
        for d in docs:
            d.metadata = {**(d.metadata or {}), **_meta(path)}
        return docs
    if suffix == ".json":
        return _load_json(path)

    docs = TextLoader(str(path), encoding="utf-8").load()
    for d in docs:
        d.metadata = {**(d.metadata or {}), **_meta(path)}
    return docs


def load_raw_documents(raw_dir: Path | None = None) -> list[Document]:
    raw_dir = raw_dir or config.KNOWLEDGE_RAW_DIR
    if not raw_dir.exists():
        raise FileNotFoundError(f"语料目录不存在：{raw_dir}")

    files = sorted(
        p for p in raw_dir.rglob("*") if p.is_file() and p.suffix.lower() in _SUPPORTED
    )
    if not files:
        raise FileNotFoundError(f"未在 {raw_dir} 找到支持的文档：{_SUPPORTED}")

    docs: list[Document] = []
    for path in files:
        part = _load_one(path)
        print(f"  loaded {path.name}: {len(part)} doc(s)")
        docs.extend(part)
    return docs
