# -*- coding: utf-8 -*-
"""内容六维自测：canon 完整性 / 一致性 + raw 派生保真。"""
from __future__ import annotations

import csv
import json
import re
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent
CANON = ROOT / "canon"
RAW = ROOT / "raw"

ISSUES: list[tuple[str, str, str]] = []  # dim, severity, msg
OKS: list[str] = []


def add(dim: str, severity: str, msg: str) -> None:
    ISSUES.append((dim, severity, msg))


def ok(msg: str) -> None:
    OKS.append(msg)


def main() -> None:
    canon_md = {p.name: p.read_text(encoding="utf-8") for p in sorted(CANON.glob("[0-9]*.md"))}
    faq_path = CANON / "10_典型客服问答.json"
    faq = json.loads(faq_path.read_text(encoding="utf-8"))
    numbers_text = (CANON / "_numbers.md").read_text(encoding="utf-8")
    all_canon = "\n".join(canon_md.values()) + "\n" + faq_path.read_text(encoding="utf-8") + "\n" + numbers_text

    # ---------- D1 结构 ----------
    expected = [
        "00_文档说明与适用范围.md",
        "01_产品简介.md",
        "02_套餐与价格.md",
        "03_额度与超额计费.md",
        "04_成员权限与协作.md",
        "05_发票与抬头.md",
        "06_退款规则.md",
        "07_数据保留与注销.md",
        "08_企业版支持与SLA.md",
        "09_工单与服务渠道说明.md",
        "10_典型客服问答.json",
        "_numbers.md",
    ]
    for name in expected:
        if not (CANON / name).exists():
            add("D1", "P0", f"缺失文件 {name}")
        else:
            ok(f"D1 存在 {name}")

    sections = [
        "## 适用范围",
        "## 规则正文",
        "## 例外与边界",
        "## 客服话术要点",
        "## 禁止承诺",
        "## 相关主题",
    ]
    for name, text in canon_md.items():
        miss = [s for s in sections if s not in text]
        if miss:
            add("D1", "P1", f"{name} 缺章节: {miss}")
        else:
            ok(f"D1 {name} 章节齐全")

    if not (30 <= len(faq) <= 40):
        add("D1", "P1", f"FAQ 条数 {len(faq)} 不在 30～40")
    else:
        ok(f"D1 FAQ {len(faq)} 条")

    # 必答覆盖（答案侧关键词）
    req_ans = [
        ("产品定位", r"AI 知识管理|智能问答 SaaS"),
        ("试用14天", r"14\s*天"),
        ("基础价99", r"99\s*元"),
        ("私有化非默认", r"单独签约"),
        ("API10000", r"10000"),
        ("额度分开", r"互不挪用|分开统计"),
        ("成员上限10", r"基础版正式成员上限为 10|成员上限[：:]*\s*\*\*10\*\*|正式成员上限[^\n]*10"),
        ("协作者不占", r"不占用正式成员"),
        ("审计不可删", r"不能删除已落库|清空审计"),
        ("支持专票", r"增值税专用发票|专票"),
        ("抬头/红字", r"红字"),
        ("全额退10%", r"10%"),
        ("拒退50%", r"50%"),
        ("到账3-7", r"3～7|3到7"),
        ("试用不立刻删", r"不会立刻"),
        ("注销", r"注销"),
        ("CSM", r"CSM|客户成功"),
        ("工单1工作日", r"1\s*个工作日"),
        ("SLA合同优先", r"以签署合同为准|不能代替合同"),
        ("ORD走工具", r"订单查询工具|业务查询|lookup_order|禁止.*编造"),
    ]
    for label, pat in req_ans:
        if re.search(pat, all_canon):
            ok(f"D1/D4 必答点命中: {label}")
        else:
            add("D1", "P0", f"必答点未定位: {label}")

    # ---------- D2 数字一致性 ----------
    # 从 _numbers 抽关键值，在正文中核对
    must_same = [
        ("99", "基础版月价"),
        ("199", "专业版月价"),
        ("10000", "基础版 API"),
        ("80000", "专业版 API"),
        ("14 天", "试用天数写法"),
    ]
    for token, label in must_same:
        in_num = token.replace(" ", "") in numbers_text.replace(" ", "") or token in numbers_text
        in_body = token in all_canon
        if not in_num:
            add("D2", "P1", f"_numbers 缺少 {label} ({token})")
        if not in_body:
            add("D2", "P0", f"正文缺少 {label} ({token})")
        if in_num and in_body:
            ok(f"D2 {label} 出现")

    # FAQ 与正文冲突探针
    faq_blob = json.dumps(faq, ensure_ascii=False)
    if "10 人" not in faq_blob and "上限为 10" not in faq_blob and "上限 10" not in faq_blob:
        add("D2", "P0", "FAQ 未明确基础版 10 人")
    else:
        ok("D2 FAQ 含基础版 10 人")

    # 危险不一致：若出现「立刻删除」且无「不会立刻」上下文 —— 已有不会立刻则 OK
    if re.search(r"会立刻永久删除|到期后立刻删除", all_canon) and "不会立刻" not in all_canon:
        add("D2", "P0", "存在「立刻删除」且无否定表述")
    else:
        ok("D2 试用删除口径含「不会立刻」")

    # ---------- D3 规则自洽 ----------
    if not re.search(r"不会立刻停用|不会停用", all_canon):
        add("D3", "P0", "缺少 API 超额不停用口径")
    else:
        ok("D3 API 超额不停用")

    if not re.search(r"停止继续提供问答|停止.*问答服务", all_canon):
        add("D3", "P0", "缺少 AI 用尽停用口径")
    else:
        ok("D3 AI 用尽停问答")

    if re.search(r"API[^\n]{0,40}立刻停用", all_canon):
        add("D3", "P0", "出现 API 立刻停用，与主规则冲突")

    if "先完成红字" not in all_canon and "先红字" not in all_canon:
        add("D3", "P1", "专票与退款红字顺序未写清")
    else:
        ok("D3 专票后退款先红字")

    if "不占用正式成员" not in all_canon:
        add("D3", "P0", "外部协作者不占名额未写清")
    else:
        ok("D3 协作者不占名额")

    # 通道边界
    if not re.search(r"ORD|订单号|业务查询|lookup_order", all_canon):
        add("D3", "P1", "未强调订单实时状态走工具")
    else:
        ok("D3 订单走工具边界有写")

    # 退款看 AI 用量而非 API
    if "AI 问答" in all_canon and ("退款" in all_canon):
        ok("D3 退款与 AI 用量有关联叙述")
    if re.search(r"API[^\n]{0,20}使用量[^\n]{0,20}退款比例", all_canon):
        add("D3", "P1", "可能把退款比例绑到 API 用量（需人工确认）")

    # ---------- D4 覆盖充分性 ----------
    themes = {
        "套餐四档": r"试用版|基础版|专业版|企业版",
        "三类额度": r"AI 问答|API|OCR",
        "成员权限": r"所有者|管理员|外部协作者",
        "发票": r"专票|普通发票",
        "退款": r"退款",
        "数据保留": r"冻结|宽限期|回收站",
        "企业SLA": r"SLA|CSM",
        "工单渠道": r"工单",
        "禁止承诺": r"禁止承诺",
    }
    for label, pat in themes.items():
        if re.search(pat, all_canon):
            ok(f"D4 覆盖 {label}")
        else:
            add("D4", "P0", f"主题覆盖不足: {label}")

    # 字数粗算
    cjk = len(re.findall(r"[\u4e00-\u9fff]", "\n".join(canon_md.values())))
    if cjk < 12000:
        add("D4", "P2", f"正文汉字约 {cjk}，略低于 A 档 1.5 万目标（可接受但偏薄）")
    else:
        ok(f"D4 正文汉字约 {cjk}")

    # ---------- D5 品牌 ----------
    bad_brands = ["atguigu", "尚硅谷", "Atguigu Assistant"]
    scan_files = list(CANON.rglob("*")) + list(RAW.rglob("*"))
    if (ROOT / "knowledge.txt").exists():
        scan_files.append(ROOT / "knowledge.txt")
    for bad in bad_brands:
        hit = []
        for p in scan_files:
            if not p.is_file():
                continue
            if p.suffix.lower() not in {".md", ".txt", ".html", ".json", ".csv", ".py"}:
                continue
            t = p.read_text(encoding="utf-8", errors="ignore")
            if bad.lower() in t.lower():
                hit.append(str(p.relative_to(ROOT)))
        if hit:
            add("D5", "P0", f"残留品牌 {bad}: {hit}")
        else:
            ok(f"D5 无 {bad}")

    if "智答云" not in all_canon:
        add("D5", "P0", "canon 未统一产品名「智答云」")
    else:
        ok("D5 产品名智答云")

    # ---------- D6 raw 保真 ----------
    mapping = {
        "md": [
            "00_文档说明与适用范围.md",
            "01_产品简介.md",
            "03_额度与超额计费.md",
            "07_数据保留与注销.md",
        ],
        "txt": ["02_套餐与价格.txt"],
        "html": ["04_成员权限与协作.html"],
        "csv": ["05_发票规则.csv"],
        "pdf": ["06_退款规则.pdf"],
        "docx": ["08_企业版支持与SLA.docx", "09_工单与服务渠道说明.docx"],
        "json": ["10_典型客服问答.json"],
    }
    for sub, files in mapping.items():
        for fn in files:
            p = RAW / sub / fn
            if not p.exists() or p.stat().st_size == 0:
                add("D6", "P0", f"raw 缺失或空: {sub}/{fn}")
            else:
                ok(f"D6 存在 {sub}/{fn} ({p.stat().st_size}B)")

    # 旧文件残留（教程短文名）
    legacy = list(RAW.rglob("*atguigu*")) + list(RAW.rglob("*尚硅谷*"))
    # 旧编号若仍在且内容短？
    old_names = [
        "02_套餐说明.txt",
        "00_文档说明.md",
        "01_产品简介.md",  # 仍存在但应是新内容
    ]
    # 检查 raw 是否还有未映射的杂文件
    expected_files = {f for files in mapping.values() for f in files}
    extras = []
    for p in RAW.rglob("*"):
        if p.is_file() and p.suffix.lower() in {
            ".md",
            ".txt",
            ".html",
            ".csv",
            ".json",
            ".pdf",
            ".docx",
        }:
            if p.name.endswith(".tmp"):
                add("D6", "P2", f"残留临时文件 {p.name}")
            elif p.name not in expected_files:
                extras.append(str(p.relative_to(RAW)))
    if extras:
        add("D6", "P1", f"raw 存在未在映射表中的文件: {extras}")
    else:
        ok("D6 raw 文件集合与映射一致")

    # 语义抽检
    txt02 = (RAW / "txt" / "02_套餐与价格.txt").read_text(encoding="utf-8")
    for token in ["99", "199", "10", "50", "14"]:
        if token not in txt02:
            add("D6", "P0", f"txt 套餐缺少关键数字 {token}")
    else:
        ok("D6 txt 套餐关键数字齐全")

    html04 = (RAW / "html" / "04_成员权限与协作.html").read_text(encoding="utf-8")
    if "不占用" not in html04:
        add("D6", "P0", "html 成员文档缺少「不占用」")
    else:
        ok("D6 html 含协作者不占用")

    csv05 = (RAW / "csv" / "05_发票规则.csv").read_text(encoding="utf-8-sig")
    if "信息技术服务费" not in csv05 or "专票" not in csv05:
        add("D6", "P0", "csv 发票关键字段缺失")
    else:
        ok("D6 csv 发票关键字段齐全")

    pdf_text = "".join(
        (p.extract_text() or "") for p in PdfReader(str(RAW / "pdf" / "06_退款规则.pdf")).pages
    )
    pdf_cjk = len(re.findall(r"[\u4e00-\u9fff]", pdf_text))
    if pdf_cjk < 200:
        add("D6", "P0", f"PDF 可抽汉字过少: {pdf_cjk}（疑似空/字体问题）")
    else:
        ok(f"D6 PDF 可抽汉字 {pdf_cjk}")
    for token in ["自然日", "50", "红字", "企业版"]:
        if token not in pdf_text:
            add("D6", "P1", f"PDF 文本缺少「{token}」")
    if all(t in pdf_text for t in ["自然日", "50", "红字"]):
        ok("D6 PDF 退款关键语义在")

    doc08 = "\n".join(
        p.text for p in DocxDocument(str(RAW / "docx" / "08_企业版支持与SLA.docx")).paragraphs
    )
    doc09 = "\n".join(
        p.text for p in DocxDocument(str(RAW / "docx" / "09_工单与服务渠道说明.docx")).paragraphs
    )
    if "私有化" not in doc08 or "合同" not in doc08:
        add("D6", "P0", "docx08 缺少私有化/合同口径")
    else:
        ok("D6 docx08 企业口径在")
    if "工单" not in doc09:
        add("D6", "P0", "docx09 缺少工单")
    else:
        ok("D6 docx09 工单口径在")

    raw_faq = json.loads((RAW / "json" / "10_典型客服问答.json").read_text(encoding="utf-8"))
    if raw_faq != faq:
        add("D6", "P0", "raw FAQ 与 canon FAQ 不一致")
    else:
        ok("D6 FAQ canon=raw")

    # md 与 canon 是否逐字一致
    for fn in mapping["md"]:
        c = (CANON / fn).read_text(encoding="utf-8")
        r = (RAW / "md" / fn).read_text(encoding="utf-8")
        if c != r:
            add("D6", "P0", f"md 与 canon 不一致: {fn}")
        else:
            ok(f"D6 md 同步 {fn}")

    kt_path = ROOT / "knowledge.txt"
    if not kt_path.exists():
        add("D6", "P0", "knowledge.txt 不存在（T1 入口说明丢失）")
    else:
        kt = kt_path.read_text(encoding="utf-8")
        if "真源" not in kt and "不再" not in kt:
            add("D6", "P1", "knowledge.txt 未体现 T1 入口说明")
        if "99 元 / 用户" in kt and "试用版" in kt and len(kt) > 2000:
            add("D6", "P0", "knowledge.txt 仍像完整第二真源")
        else:
            ok("D6 knowledge.txt 已是短入口（T1）")

    # ---------- 报告 ----------
    print("=" * 60)
    print("知识库内容六维自测报告")
    print("=" * 60)
    print(f"通过项: {len(OKS)}")
    p0 = [i for i in ISSUES if i[1] == "P0"]
    p1 = [i for i in ISSUES if i[1] == "P1"]
    p2 = [i for i in ISSUES if i[1] == "P2"]
    print(f"问题: P0={len(p0)} P1={len(p1)} P2={len(p2)}")
    for sev in ("P0", "P1", "P2"):
        items = [i for i in ISSUES if i[1] == sev]
        if not items:
            continue
        print(f"\n--- {sev} ---")
        for dim, _, msg in items:
            print(f"[{dim}] {msg}")

    if not p0 and not p1:
        verdict = "PASS（可进入 ingest/手测）"
    elif not p0:
        verdict = "PASS WITH NOTES（有 P1/P2，建议修但不挡派生）"
    else:
        verdict = "FAIL（存在 P0，建议先修再入库）"
    print("\n结论:", verdict)


if __name__ == "__main__":
    main()
