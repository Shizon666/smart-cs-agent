# -*- coding: utf-8 -*-
"""从 knowledge/canon 派生 knowledge/raw 多格式语料（Phase 2）。"""
from __future__ import annotations

import csv
import re
import shutil
from pathlib import Path

from docx import Document as DocxDocument
from fpdf import FPDF

ROOT = Path(__file__).resolve().parent.parent
CANON = ROOT / "knowledge" / "canon"
RAW = ROOT / "knowledge" / "raw"

# Windows 中文字体：优先 TTF（TTC 在 fpdf2 上常「子集失败 → 预览空白」）
_FONT_CANDIDATES = [
    Path(r"C:\Windows\Fonts\simhei.ttf"),
    Path(r"C:\Windows\Fonts\simkai.ttf"),
    Path(r"C:\Windows\Fonts\msyh.ttf"),
    Path(r"C:\Windows\Fonts\simsun.ttc"),
    Path(r"C:\Windows\Fonts\msyh.ttc"),
]


def _read(name: str) -> str:
    return (CANON / name).read_text(encoding="utf-8")


def _strip_md(text: str) -> str:
    text = re.sub(r"^>\s*", "", text, flags=re.M)
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"^#+\s*", "", text, flags=re.M)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text.strip() + "\n"


def _clear_raw() -> None:
    """确保子目录存在；不整目录删除，避免 Windows 下 PDF 被占用时整次失败。"""
    for sub in ("md", "txt", "html", "csv", "json", "pdf", "docx"):
        (RAW / sub).mkdir(parents=True, exist_ok=True)


def _write_md() -> None:
    for name in (
        "00_文档说明与适用范围.md",
        "01_产品简介.md",
        "03_额度与超额计费.md",
        "07_数据保留与注销.md",
    ):
        (RAW / "md" / name).write_text(_read(name), encoding="utf-8")
        print("md", name)


def _write_txt() -> None:
    name = "02_套餐与价格.md"
    out = RAW / "txt" / "02_套餐与价格.txt"
    out.write_text(_strip_md(_read(name)), encoding="utf-8")
    print("txt", out.name)


def _write_html() -> None:
    body = _strip_md(_read("04_成员权限与协作.md"))
    # 简单换行转段落
    paras = "".join(f"<p>{line}</p>\n" if line.strip() else "<br/>\n" for line in body.splitlines())
    html = (
        "<!DOCTYPE html>\n<html lang=\"zh-CN\">\n<head>"
        "<meta charset=\"utf-8\"/>"
        "<title>智答云 · 成员权限与协作</title>"
        "</head>\n<body>\n"
        f"{paras}"
        "</body>\n</html>\n"
    )
    out = RAW / "html" / "04_成员权限与协作.html"
    out.write_text(html, encoding="utf-8")
    print("html", out.name)


def _write_csv() -> None:
    """DOC05 全程口径落 CSV（仍只此一种 raw 格式，行级覆盖全文要点）。"""
    rows = [
        ["条款", "规则项", "口径", "备注"],
        # 元信息
        ["DOC05", "文档说明", "智答云发票与抬头规则；知识库版本2026-Q3；生效2026-07-01", "中国区订阅"],
        ["DOC05", "适用范围", "电子普票与增值税专票、开票时间、金额口径、抬头错误处理；退款与专票红字交叉见DOC06", ""],
        # 5.1
        ["5.1", "发票类型", "支持电子普通发票、增值税专用发票（专票）", ""],
        ["5.1", "默认开票内容", "信息技术服务费", ""],
        ["5.1", "特殊开票内容", "须在付款前联系销售确认是否可开；付款后无法保证变更成功", "禁止事后随便改类目"],
        # 5.2
        ["5.2", "月付开票时机", "支付成功后即可申请开票", ""],
        ["5.2", "退款与开票", "同一自然月内发生退款，退款对应部分不能重复开票", ""],
        ["5.2", "企业年付开票", "款项到账并完成合同归档后统一开票", ""],
        ["5.2", "开具时效", "资料齐全时电子票通常申请后1～3个工作日内开具；高峰可能延长，以工单反馈为准", ""],
        # 5.3
        ["5.3", "发票金额", "按实际支付金额开具，不含已退还部分", ""],
        ["5.3", "优惠券", "券抵扣金额不开票，只对用户实付部分开票", ""],
        # 5.4
        ["5.4", "专票资料", "须维护公司名称、纳税人识别号、地址电话、开户行及账号等（以开票页必填为准）；信息不全驳回", ""],
        # 5.5 抬头（完整分步）
        ["5.5", "抬头错误-尚未开具", "可在开票申请页自行修改抬头与税号后重新提交", ""],
        ["5.5", "抬头错误-已开普票未报销", "可提交工单申请作废或红冲后重开，须提供正确抬头；是否允许以财务审核为准；同一订单原则上限制次数", ""],
        ["5.5", "抬头错误-已开专票", "必须走红字发票信息确认流程后再重开正确专票", ""],
        ["5.5", "已开专票且要退款", "必须先完成红字流程再退款（见DOC06）", "客服勿说马上退"],
        ["5.5", "多次重开", "用户自身填错导致多次重开，可能要求加盖公章说明；不得承诺无限次免费重开", ""],
        # 5.6 速查
        ["5.6", "是否支持专票", "支持", ""],
        ["5.6", "专票后退款顺序", "先红字，后退款", ""],
        # 5.7
        ["5.7", "公司更名", "提供工商变更证明与新开票资料，经财务审核后用于后续订单；已开旧票是否换开按税务与财务审核，不保证所有历史票都能换开", ""],
        ["5.7", "合并开票", "同一付款主体多笔月付是否合并一张，以财务当期能力为准；不能承诺跨年合并", ""],
        # 5.8 误解
        ["5.8", "误解-自动开票", "付了钱不会自动开票，需用户发起申请；企业年付可能由销售统一申请", ""],
        ["5.8", "误解-退款后仍报销原票", "退款对应部分不可重复作为有效应税凭证；专票须红字", ""],
        ["5.8", "误解-个人开专票", "个人抬头通常不符合专票条件，应开普票", ""],
        # 5.9 清单
        ["5.9", "开票检查清单1", "确认订单已支付成功且未全额退款", ""],
        ["5.9", "开票检查清单2", "选择普票或专票；专票备齐名称税号地址电话开户行账号", ""],
        ["5.9", "开票检查清单3", "开票内容默认信息技术服务费；特殊内容须付款前已确认", ""],
        ["5.9", "开票检查清单4", "核对金额等于实付（不含券）", ""],
        ["5.9", "开票检查清单5", "提交后保存申请单号便于工单追踪", ""],
        # 例外/话术/禁诺
        ["例外", "专票资质", "个人是否具备专票资质以税务与平台审核为准；不符合则开普票", ""],
        ["例外", "复杂主体", "第三方代付、加盟商代开等升级人工财务处理", ""],
        ["例外", "下载链接过期", "电子票下载链接过期可工单重发，不视为重新开具金额", ""],
        ["话术", "开场三问", "先问：普票还是专票、是否已开出、是否已报销", ""],
        ["话术", "退款提醒", "有专票别直接说马上退，先讲红字", ""],
        ["禁止", "开票内容", "禁止承诺法定以外或未确认的开票内容（如随意改成软件销售）", ""],
        ["禁止", "当日必开", "资料不齐时禁止承诺今天一定开出", ""],
        ["禁止", "税率鉴定", "禁止承诺一定免税或具体税率解释；税率以实际开具为准", ""],
        ["关联", "相关主题", "DOC06退款与红字；DOC02支付；DOC09发票类工单", ""],
    ]
    out = RAW / "csv" / "05_发票规则.csv"
    with out.open("w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f).writerows(rows)
    print("csv", out.name, "rows=", len(rows) - 1)


def _write_json() -> None:
    src = CANON / "10_典型客服问答.json"
    dst = RAW / "json" / "10_典型客服问答.json"
    shutil.copy2(src, dst)
    print("json", dst.name)


def _font_path() -> Path:
    for p in _FONT_CANDIDATES:
        if p.exists():
            return p
    raise FileNotFoundError("未找到中文字体（simhei/msyh/simsun），无法导出 PDF")


def _write_pdf() -> None:
    text = _strip_md(_read("06_退款规则.md"))
    text = text.replace("|", " ")
    text = re.sub(r"-{3,}", "——", text)
    font = _font_path()
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    pdf.set_margins(18, 18, 18)
    pdf.add_font("cn", fname=str(font))
    pdf.set_font("cn", size=11)
    usable = pdf.w - pdf.l_margin - pdf.r_margin
    for line in text.splitlines():
        content = line.strip() if line.strip() else " "
        pdf.multi_cell(usable, 7, content)
    out = RAW / "pdf" / "06_退款规则.pdf"
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_suffix(".pdf.tmp")
    pdf.output(str(tmp))
    try:
        if out.exists():
            out.unlink()
        tmp.replace(out)
        final = out
    except PermissionError:
        # 常见：资源管理器/预览占用 PDF
        final = tmp
        print(
            "WARN: 无法覆盖 06_退款规则.pdf（文件被占用）。"
            f"已写入 {tmp.name}，请关闭预览后重跑或手动改名。"
        )

    from pypdf import PdfReader

    extracted = "".join((p.extract_text() or "") for p in PdfReader(str(final)).pages)
    cjk = len(re.findall(r"[\u4e00-\u9fff]", extracted))
    if cjk < 200:
        raise RuntimeError(f"PDF 导出异常：可抽取汉字仅 {cjk}，请检查字体 {font}")
    print("pdf", final.name, "font=", font.name, "cjk=", cjk, "bytes=", final.stat().st_size)


def _md_to_docx(md_name: str, docx_name: str) -> None:
    text = _strip_md(_read(md_name))
    doc = DocxDocument()
    doc.add_heading("智答云知识库", level=0)
    for line in text.splitlines():
        doc.add_paragraph(line if line.strip() else "")
    out = RAW / "docx" / docx_name
    doc.save(str(out))
    print("docx", out.name)


def _write_knowledge_txt_t1() -> None:
    path = ROOT / "knowledge" / "knowledge.txt"
    path.write_text(
        "智答云客服知识库（入口说明 · 2026-Q3）\n"
        "\n"
        "本文件不再作为业务口径真源，也不应在 raw/ 非空时被 ingest 读取。\n"
        "请使用：\n"
        "  - knowledge/canon/     唯一真源（撰写与改数）\n"
        "  - knowledge/raw/       ingest 扫描入口（由 canon 派生）\n"
        "\n"
        "派生命令：uv run python -m knowledge.export_raw\n"
        "入库命令：uv run python -m rag.ingest --recreate\n"
        "设计文档：docs/smart-cs-agent-业务知识库与RAG调优设计文档.md\n",
        encoding="utf-8",
    )
    print("knowledge.txt T1 updated")


def append_canon_extras() -> None:
    """补一点字数，贴近 A 档。"""
    extras = {
        "04_成员权限与协作.md": """
### 4.9 角色对照速查（可检索）

| 动作 | 所有者 | 管理员 | 普通成员 | 外部协作者 |
|------|--------|--------|----------|------------|
| 上传授权库文档 | 是 | 是 | 视授权 | 视授权 |
| 邀请正式成员 | 是 | 是 | 否 | 否 |
| 修改计费/套餐 | 是 | 否 | 否 | 否 |
| 删除工作区 | 是 | 否 | 否 | 否 |
| 导出审计日志 | 专业版+有权限角色 | 专业版管理员通常可 | 否 | 否 |
| 清空审计日志 | 否 | 否 | 否 | 否 |
| 访问计费页 | 是 | 否 | 否 | 否 |

说明：导出审计不等于删除审计；「清空审计」全角色默认不可（见 §4.3）。
""",
        "05_发票与抬头.md": """
### 5.9 申请开票检查清单

1. 确认订单已支付成功且未全额退款；
2. 选择普票或专票；专票资料：名称、税号、地址电话、开户行账号；
3. 开票内容默认「信息技术服务费」，特殊内容须付款前已确认；
4. 核对金额=实付（不含券）；
5. 提交后保存申请单号，便于工单追踪。
""",
        "08_企业版支持与SLA.md": """
### 8.8 企业版沟通检查清单

1. 是否已签约企业版主合同？
2. 合同是否含 SLA 附件？若无，仅能引用默认响应目标，不能谈赔偿数字。
3. 是否单独签署私有化/定制附表？未签则不能承诺私有化交付。
4. CSM 与专属群是否已开通？未开通转销售。
5. 用户要合同级分钟数 → 升级人工调阅合同，禁止用本库杜撰。
""",
        "09_工单与服务渠道说明.md": """
### 9.9 渠道选择速查

| 诉求 | 优先渠道 |
|------|----------|
| 规则咨询（退款/额度/发票政策） | 智能客服知识库 + 必要时工单 |
| 实时订单/工单进度 | 业务查询工具（要单号） |
| 账号打不开/大面积故障 | 紧急工单；专业版可电话；企业版专属群 |
| 特批退款/特批留数 | 人工主管书面审批，智能客服不承诺 |
| 企业 SLA 赔偿 | 合同 + CSM/法务，不用本库数字替代 |
""",
    }
    for name, block in extras.items():
        path = CANON / name
        text = path.read_text(encoding="utf-8")
        marker = "## 例外与边界"
        key = block.strip().splitlines()[0]
        if key in text:
            print("skip extra", name)
            continue
        path.write_text(text.replace(marker, block.strip() + "\n\n" + marker, 1), encoding="utf-8")
        print("extra", name)


def main() -> None:
    append_canon_extras()
    _clear_raw()
    _write_md()
    _write_txt()
    _write_html()
    _write_csv()
    _write_json()
    _write_pdf()
    _md_to_docx("08_企业版支持与SLA.md", "08_企业版支持与SLA.docx")
    _md_to_docx("09_工单与服务渠道说明.md", "09_工单与服务渠道说明.docx")
    _write_knowledge_txt_t1()
    print("done →", RAW)


if __name__ == "__main__":
    main()
